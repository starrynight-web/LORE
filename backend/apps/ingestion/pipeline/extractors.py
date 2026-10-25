import io
import csv
import os
import re
import zipfile
import xml.etree.ElementTree as ET
from typing import Dict, Any, List, Optional
from .ocr import OCRRouter, BaseOCRProvider

class DocxExtractor:
    """Extracts structured text, headings, and tables from DOCX documents."""
    @staticmethod
    def extract(file_path_or_bytes) -> List[Dict[str, Any]]:
        import docx
        if isinstance(file_path_or_bytes, (bytes, bytearray)):
            doc = docx.Document(io.BytesIO(file_path_or_bytes))
        else:
            doc = docx.Document(file_path_or_bytes)

        lines = []
        # Extract paragraphs with heading markers
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            style_name = getattr(p.style, 'name', '') or ''
            if style_name.startswith('Heading 1'):
                lines.append(f"# {text}")
            elif style_name.startswith('Heading 2'):
                lines.append(f"## {text}")
            elif style_name.startswith('Heading 3'):
                lines.append(f"### {text}")
            else:
                lines.append(text)

        # Extract tables
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
                table_rows.append(" | ".join(cells))
            if table_rows:
                lines.append("\n" + "\n".join(table_rows) + "\n")

        full_text = "\n\n".join(lines)
        return [{
            'page_number': 1,
            'text': full_text,
            'confidence': 1.0,
            'words_count': len(full_text.split()),
            'width': 612.0,
            'height': 792.0,
            'extraction_type': 'docx_extract',
            'engine': 'python-docx',
            'engine_version': '1.2.0',
            'is_repaired': False,
            'bounding_boxes': [],
        }]

class SpreadsheetExtractor:
    """Extracts tables, rows, and columns from CSV and XLSX documents."""
    @staticmethod
    def extract_csv(file_path_or_bytes) -> List[Dict[str, Any]]:
        if isinstance(file_path_or_bytes, (bytes, bytearray)):
            text = file_path_or_bytes.decode('utf-8', errors='replace')
            reader = csv.reader(io.StringIO(text))
            rows = list(reader)
        else:
            with open(file_path_or_bytes, 'r', encoding='utf-8', errors='replace') as f:
                reader = csv.reader(f)
                rows = list(reader)

        formatted_rows = [" | ".join(r) for r in rows if any(r)]
        full_text = "\n".join(formatted_rows)
        return [{
            'page_number': 1,
            'text': full_text,
            'confidence': 1.0,
            'words_count': len(full_text.split()),
            'width': 800.0,
            'height': 600.0,
            'extraction_type': 'csv_extract',
            'engine': 'csv_native',
            'engine_version': '1.0',
            'is_repaired': False,
            'bounding_boxes': [],
        }]

    @staticmethod
    def extract_xlsx(file_path_or_bytes) -> List[Dict[str, Any]]:
        import openpyxl
        try:
            if isinstance(file_path_or_bytes, (bytes, bytearray)):
                wb = openpyxl.load_workbook(io.BytesIO(file_path_or_bytes), data_only=True)
            else:
                wb = openpyxl.load_workbook(file_path_or_bytes, data_only=True)

            pages = []
            for idx, sheet_name in enumerate(wb.sheetnames):
                ws = wb[sheet_name]
                rows_data = []
                for row in ws.iter_rows(values_only=True):
                    if any(row):
                        str_cells = [str(c) if c is not None else "" for c in row]
                        rows_data.append(" | ".join(str_cells))
                text = f"## Sheet: {sheet_name}\n\n" + "\n".join(rows_data)
                pages.append({
                    'page_number': idx + 1,
                    'text': text,
                    'confidence': 1.0,
                    'words_count': len(text.split()),
                    'width': 800.0,
                    'height': 600.0,
                    'extraction_type': 'xlsx_extract',
                    'engine': 'openpyxl',
                    'engine_version': '3.1.5',
                    'is_repaired': False,
                    'bounding_boxes': [],
                })
            return pages if pages else TextExtractor.extract(file_path_or_bytes)
        except Exception:
            return TextExtractor.extract(file_path_or_bytes)

class PresentationExtractor:
    """
    Extracts structured presentation content from PPTX files:
    - Slide order and 1-indexed numbers (page_number).
    - Slide titles and text boxes with paragraph bullet hierarchies.
    - Tables formatted into clean Markdown tables.
    - Speaker notes extracted from linked notesSlide relationships.
    - Hyperlinks and presentation Dublin Core metadata.
    - Empty and corrupted presentation handling without throwing unhandled crashes.
    """
    @staticmethod
    def _parse_rels(zf: zipfile.ZipFile, rels_path: str) -> Dict[str, Dict[str, str]]:
        rels = {}
        if rels_path not in zf.namelist():
            return rels
        try:
            tree = ET.fromstring(zf.read(rels_path))
            for node in tree:
                if node.tag.endswith('}Relationship'):
                    r_id = node.attrib.get('Id')
                    target = node.attrib.get('Target', '')
                    r_type = node.attrib.get('Type', '')
                    if r_id:
                        rels[r_id] = {'target': target, 'type': r_type}
        except Exception:
            pass
        return rels

    @staticmethod
    def _parse_metadata(zf: zipfile.ZipFile) -> Dict[str, str]:
        meta = {}
        if 'docProps/core.xml' in zf.namelist():
            try:
                tree = ET.fromstring(zf.read('docProps/core.xml'))
                for node in tree.iter():
                    tag = node.tag.split('}')[-1] if '}' in node.tag else node.tag
                    if tag in ('title', 'creator', 'lastModifiedBy', 'subject', 'description') and node.text:
                        meta[tag] = node.text.strip()
            except Exception:
                pass
        return meta

    @classmethod
    def _resolve_slide_order(cls, zf: zipfile.ZipFile) -> List[str]:
        ordered = []
        if 'ppt/presentation.xml' in zf.namelist() and 'ppt/_rels/presentation.xml.rels' in zf.namelist():
            try:
                pres_tree = ET.fromstring(zf.read('ppt/presentation.xml'))
                pres_rels = cls._parse_rels(zf, 'ppt/_rels/presentation.xml.rels')

                for node in pres_tree.iter():
                    if node.tag.endswith('}sldId'):
                        r_id = node.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id') or node.attrib.get('r:id') or node.attrib.get('id')
                        if r_id and r_id in pres_rels:
                            target = pres_rels[r_id]['target'].replace('\\', '/')
                            target = target.lstrip('/')
                            if not target.startswith('ppt/'):
                                target = f"ppt/{target}"
                            if target in zf.namelist():
                                ordered.append(target)
            except Exception:
                pass

        if not ordered:
            # Fallback to sorting filenames numerically
            slide_names = [n for n in zf.namelist() if re.match(r'ppt/slides/slide\d+\.xml', n)]
            def _get_num(s: str) -> int:
                m = re.search(r'\d+', s)
                return int(m.group(0)) if m else 0
            slide_names.sort(key=_get_num)
            ordered = slide_names

        return ordered

    @classmethod
    def _extract_slide_notes(cls, zf: zipfile.ZipFile, slide_name: str, slide_rels: Dict[str, Dict[str, str]]) -> str:
        notes_target = None
        for rel in slide_rels.values():
            if rel['type'].endswith('notesSlide'):
                notes_target = rel['target'].replace('\\', '/')
                break

        if not notes_target:
            return ''

        # Normalize relative path: e.g. ../notesSlides/notesSlide1.xml -> ppt/notesSlides/notesSlide1.xml
        base_dir = os.path.dirname(slide_name)
        norm_notes = os.path.normpath(os.path.join(base_dir, notes_target)).replace('\\', '/')
        if norm_notes not in zf.namelist():
            return ''

        try:
            tree = ET.fromstring(zf.read(norm_notes))
            paragraphs = []
            for sp in tree.iter():
                if sp.tag.endswith('}sp'):
                    # Skip slide number and footer shapes
                    is_header_footer = False
                    for ph in sp.iter():
                        if ph.tag.endswith('}ph'):
                            ph_type = ph.attrib.get('type', '')
                            if ph_type in ('sldNum', 'hdr', 'ftr', 'dt'):
                                is_header_footer = True
                                break
                    if is_header_footer:
                        continue

                    sp_texts = []
                    for t in sp.iter():
                        if t.tag.endswith('}t') and t.text:
                            sp_texts.append(t.text.strip())
                    if sp_texts:
                        paragraphs.append(" ".join(sp_texts))
            return "\n\n".join(paragraphs).strip()
        except Exception:
            return ''

    @classmethod
    def extract(cls, file_path_or_bytes) -> List[Dict[str, Any]]:
        try:
            if isinstance(file_path_or_bytes, (bytes, bytearray)):
                stream = io.BytesIO(file_path_or_bytes)
            else:
                stream = file_path_or_bytes

            if not zipfile.is_zipfile(stream):
                return TextExtractor.extract(file_path_or_bytes)

            zf = zipfile.ZipFile(stream)
            ordered_slides = cls._resolve_slide_order(zf)

            if not ordered_slides:
                return [{
                    'page_number': 1,
                    'text': '# Presentation\n\n[Empty presentation or no slides found]',
                    'confidence': 1.0,
                    'words_count': 0,
                    'width': 960.0,
                    'height': 540.0,
                    'extraction_type': 'pptx_structured_extract',
                    'engine': 'openxml_pptx',
                    'engine_version': '2.0.0',
                    'is_repaired': False,
                    'bounding_boxes': []
                }]

            meta = cls._parse_metadata(zf)
            pres_header = ""
            if meta.get('title'):
                pres_header = f"Presentation: {meta['title']}\n"
                if meta.get('creator'):
                    pres_header += f"Author: {meta['creator']}\n"

            pages = []
            for idx, sname in enumerate(ordered_slides, 1):
                try:
                    xml_data = zf.read(sname)
                    root = ET.fromstring(xml_data)
                except Exception:
                    pages.append({
                        'page_number': idx,
                        'text': f"# Slide {idx}\n\n[Corrupted slide data]",
                        'confidence': 0.0,
                        'words_count': 0,
                        'width': 960.0,
                        'height': 540.0,
                        'extraction_type': 'pptx_structured_extract',
                        'engine': 'openxml_pptx',
                        'engine_version': '2.0.0',
                        'is_repaired': False,
                        'bounding_boxes': []
                    })
                    continue

                # Load slide relationships for hyperlinks and speaker notes
                slide_dir = os.path.dirname(sname)
                slide_base = os.path.basename(sname)
                rels_path = f"{slide_dir}/_rels/{slide_base}.rels"
                slide_rels = cls._parse_rels(zf, rels_path)

                # Extract slide title
                slide_title = ""
                title_shapes = []
                for sp in root.iter():
                    if sp.tag.endswith('}sp'):
                        for ph in sp.iter():
                            if ph.tag.endswith('}ph'):
                                p_type = ph.attrib.get('type', '')
                                if p_type in ('title', 'ctrTitle', 'subTitle'):
                                    title_shapes.append(sp)
                                    break

                if title_shapes:
                    title_texts = []
                    for t in title_shapes[0].iter():
                        if t.tag.endswith('}t') and t.text:
                            title_texts.append(t.text.strip())
                    slide_title = " ".join(title_texts).strip()

                # Extract body paragraphs with bullet levels
                slide_paragraphs = []
                for sp in root.iter():
                    if sp.tag.endswith('}sp'):
                        if title_shapes and sp is title_shapes[0]:
                            continue  # Already extracted as title

                        for p in sp.iter():
                            if p.tag.endswith('}p'):
                                lvl = 0
                                for pPr in p.iter():
                                    if pPr.tag.endswith('}pPr'):
                                        lvl = int(pPr.attrib.get('lvl', 0))
                                        break

                                runs_text = []
                                for r in p.iter():
                                    if r.tag.endswith('}r'):
                                        hlink_url = ""
                                        for hlink in r.iter():
                                            if hlink.tag.endswith('}hlinkClick'):
                                                r_id = hlink.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id') or hlink.attrib.get('r:id')
                                                if r_id and r_id in slide_rels:
                                                    hlink_url = slide_rels[r_id]['target']
                                                break

                                        t_text = "".join(t.text for t in r.iter() if t.tag.endswith('}t') and t.text)
                                        if t_text:
                                            if hlink_url:
                                                runs_text.append(f"[{t_text.strip()}]({hlink_url})")
                                            else:
                                                runs_text.append(t_text)

                                full_p = "".join(runs_text).strip()
                                if full_p:
                                    indent = "  " * lvl
                                    prefix = f"{indent}- " if lvl > 0 else ""
                                    slide_paragraphs.append(f"{prefix}{full_p}")

                # Extract tables
                markdown_tables = []
                for tbl in root.iter():
                    if tbl.tag.endswith('}tbl'):
                        rows = []
                        for tr in tbl.iter():
                            if tr.tag.endswith('}tr'):
                                cell_texts = []
                                for tc in tr.iter():
                                    if tc.tag.endswith('}tc'):
                                        t_words = [t.text.strip() for t in tc.iter() if t.tag.endswith('}t') and t.text]
                                        cell_texts.append(" ".join(t_words).replace('\n', ' ').replace('|', '\\|'))
                                if any(cell_texts):
                                    rows.append(cell_texts)

                        if rows:
                            max_cols = max(len(r) for r in rows)
                            normalized_rows = [r + [''] * (max_cols - len(r)) for r in rows]
                            header = "| " + " | ".join(normalized_rows[0]) + " |"
                            separator = "| " + " | ".join(['---'] * max_cols) + " |"
                            body_lines = ["| " + " | ".join(r) + " |" for r in normalized_rows[1:]]
                            markdown_tables.append("\n".join([header, separator] + body_lines))

                # Extract slide hyperlinks
                hyperlinks = []
                for rel in slide_rels.values():
                    target = rel.get('target', '')
                    if target.startswith(('http://', 'https://', 'mailto:')):
                        if target not in hyperlinks:
                            hyperlinks.append(target)

                # Extract speaker notes
                speaker_notes = cls._extract_slide_notes(zf, sname, slide_rels)

                # Assemble slide markdown
                slide_sections = []
                if idx == 1 and pres_header:
                    slide_sections.append(f"> {pres_header.strip()}\n")

                if slide_title:
                    slide_sections.append(f"# Slide {idx}: {slide_title}")
                else:
                    slide_sections.append(f"# Slide {idx}")

                if slide_paragraphs:
                    slide_sections.append("\n".join(slide_paragraphs))

                if markdown_tables:
                    for tbl_md in markdown_tables:
                        slide_sections.append("\n" + tbl_md + "\n")

                if hyperlinks:
                    links_md = "*Hyperlinks referenced:*\n" + "\n".join(f"- <{link}>" for link in hyperlinks)
                    slide_sections.append(links_md)

                if speaker_notes:
                    slide_sections.append(f"### Speaker Notes:\n{speaker_notes}")

                if len(slide_sections) == 1:
                    slide_sections.append("[Empty slide]")

                final_text = "\n\n".join(slide_sections).strip()

                pages.append({
                    'page_number': idx,
                    'text': final_text,
                    'confidence': 1.0,
                    'words_count': len(final_text.split()),
                    'width': 960.0,
                    'height': 540.0,
                    'extraction_type': 'pptx_structured_extract',
                    'engine': 'openxml_pptx',
                    'engine_version': '2.0.0',
                    'is_repaired': False,
                    'bounding_boxes': []
                })

            return pages
        except Exception:
            return TextExtractor.extract(file_path_or_bytes)


class HTMLExtractor:
    """Extracts clean, semantic Markdown/text from HTML files, eliminating script and styles to prevent hallucination."""
    @staticmethod
    def extract(file_path_or_bytes) -> List[Dict[str, Any]]:
        if isinstance(file_path_or_bytes, (bytes, bytearray)):
            raw_html = file_path_or_bytes.decode('utf-8', errors='replace')
        else:
            with open(file_path_or_bytes, 'r', encoding='utf-8', errors='replace') as f:
                raw_html = f.read()

        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(raw_html, 'html.parser')
            for tag in soup(['script', 'style', 'head', 'meta', 'link', 'svg', 'noscript']):
                tag.decompose()
            for h in soup.find_all(['h1', 'h2', 'h3', 'h4', 'h5', 'h6']):
                level = int(h.name[1])
                h.replace_with(f"\n\n{'#' * level} {h.get_text().strip()}\n")
            text = soup.get_text(separator="\n", strip=True)
        except Exception:
            text = re.sub(r'<(script|style)[^>]*>.*?</\1>', '', raw_html, flags=re.DOTALL | re.IGNORECASE)
            text = re.sub(r'<[^>]+>', ' ', text)
            text = re.sub(r'\s+', ' ', text).strip()

        return [{
            'page_number': 1,
            'text': text.strip(),
            'confidence': 1.0,
            'words_count': len(text.split()),
            'width': 612.0,
            'height': 792.0,
            'extraction_type': 'html_semantic_extract',
            'engine': 'bs4_semantic',
            'engine_version': '1.0',
            'is_repaired': False,
            'bounding_boxes': [],
        }]

class TextExtractor:
    """Extracts plain text and markdown files."""
    @staticmethod
    def extract(file_path_or_bytes) -> List[Dict[str, Any]]:
        if isinstance(file_path_or_bytes, (bytes, bytearray)):
            text = file_path_or_bytes.decode('utf-8', errors='replace')
        else:
            with open(file_path_or_bytes, 'r', encoding='utf-8', errors='replace') as f:
                text = f.read()

        return [{
            'page_number': 1,
            'text': text.strip(),
            'confidence': 1.0,
            'words_count': len(text.split()),
            'width': 612.0,
            'height': 792.0,
            'extraction_type': 'plaintext_extract',
            'engine': 'native_utf8',
            'engine_version': '1.0',
            'is_repaired': False,
            'bounding_boxes': [],
        }]

class UnifiedDocumentExtractor:
    """Dispatches extraction based on document type."""
    def __init__(self, ocr_provider: Optional[BaseOCRProvider] = None):
        self.router = OCRRouter(ocr_provider=ocr_provider)

    def extract_document(self, file_path_or_bytes, doc_type: str, filename: str = '') -> List[Dict[str, Any]]:
        ext = os.path.splitext(filename.lower())[1] if filename else ''

        if doc_type == 'pdf' or ext == '.pdf':
            return self.router.process_pdf(file_path_or_bytes)
        elif doc_type == 'image' or ext in ('.png', '.jpg', '.jpeg', '.tiff', '.tif', '.webp'):
            return self.router.process_image(file_path_or_bytes)
        elif doc_type == 'docx' or ext == '.docx':
            return DocxExtractor.extract(file_path_or_bytes)
        elif ext == '.csv':
            return SpreadsheetExtractor.extract_csv(file_path_or_bytes)
        elif doc_type == 'spreadsheet' or ext in ('.xlsx', '.xls'):
            return SpreadsheetExtractor.extract_xlsx(file_path_or_bytes)
        elif doc_type == 'presentation' or ext in ('.pptx', '.ppt'):
            return PresentationExtractor.extract(file_path_or_bytes)
        elif ext in ('.html', '.htm'):
            return HTMLExtractor.extract(file_path_or_bytes)
        else:
            return TextExtractor.extract(file_path_or_bytes)
